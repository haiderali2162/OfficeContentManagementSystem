"""Validation and persistence for DC valuation rates and Excel imports."""

from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import csv
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import or_, select
from sqlalchemy.orm import joinedload

from core.database import SessionLocal
from core.models import DC_RATE_UNITS, DCValuationRate, PropertyLocation


class DCRateServiceError(ValueError):
    """Raised when a DC-rate operation fails validation or conflicts."""


LOCATION_FIELDS = (
    "name",
    "province",
    "district",
    "tehsil",
    "land_type",
    "committee_town",
    "revenue_circle",
    "property_area",
    "qanoongoi",
    "patwar_circle",
    "mauza",
    "hadbast_number",
    "square_no",
    "kila_no",
    "khasra_no",
    "locality",
    "notes",
)
LOCATION_IDENTITY_FIELDS = tuple(
    field for field in LOCATION_FIELDS if field != "notes"
)
RATE_HEADERS = {
    "location": "location_name",
    "location_name": "location_name",
    "province": "province",
    "district": "district",
    "tehsil": "tehsil",
    "rural_or_urban": "land_type",
    "land_type": "land_type",
    "committee": "committee_town",
    "town": "committee_town",
    "committee_town": "committee_town",
    "revenue_circle": "revenue_circle",
    "property_area": "property_area",
    "qanoongoi": "qanoongoi",
    "qanungoi": "qanoongoi",
    "patwar_circle": "patwar_circle",
    "mauza": "mauza",
    "hadbast": "hadbast_number",
    "hadbast_number": "hadbast_number",
    "square_no": "square_no",
    "square_number": "square_no",
    "kila_no": "kila_no",
    "killa_no": "kila_no",
    "khasra_no": "khasra_no",
    "khasra_number": "khasra_no",
    "locality": "locality",
    "location_description": "locality",
    "property_type": "property_type",
    "classification": "classification",
    "land_classification": "classification",
    "floor": "floor",
    "schedule_type": "schedule_type",
    "schedule": "schedule_type",
    "dc_schedule": "rate_value",
    "fbr_schedule": "rate_value",
    "schedule_reference": "source_reference",
    "rate": "rate_value",
    "rate_value": "rate_value",
    "rate_unit": "rate_unit",
    "unit": "rate_unit",
    "effective_from": "effective_from",
    "published_year": "published_year",
    "publish_year": "published_year",
    "schedule_year": "published_year",
    "year": "published_year",
    "source": "source_reference",
    "source_reference": "source_reference",
}
REQUIRED_RATE_COLUMNS = {
    "location_name",
    "property_type",
    "rate_value",
    "rate_unit",
}


def _text(value: Any) -> str | None:
    """Trim optional text, mapping empty values to None."""
    if value is None:
        return None
    normalized = str(value).strip()
    return normalized or None


def _date(value: Any, label: str, *, required: bool = False) -> date | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise DCRateServiceError(f"{label} is required.")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except (TypeError, ValueError):
        raise DCRateServiceError(f"{label} must use YYYY-MM-DD format.") from None


def _location_data(values: dict[str, Any]) -> dict[str, str | None]:
    name = _text(values.get("name", values.get("location_name")))
    if not name:
        raise DCRateServiceError("Location name is required.")
    location = {field: _text(values.get(field)) for field in LOCATION_FIELDS}
    location["name"] = name
    return location


def _location_key(values: dict[str, Any]) -> tuple[str, ...]:
    return tuple(
        (values.get(field) or "").strip().casefold()
        for field in LOCATION_IDENTITY_FIELDS
    )


def _rate_data(values: dict[str, Any]) -> dict[str, Any]:
    property_type = _text(values.get("property_type") or values.get("classification"))
    if not property_type:
        raise DCRateServiceError("Property type is required.")

    try:
        amount = Decimal(str(values.get("rate_value", "")).strip())
    except (InvalidOperation, ValueError, TypeError):
        raise DCRateServiceError("Rate must be a valid decimal number.") from None
    if not amount.is_finite() or amount < 0:
        raise DCRateServiceError("Rate must be a finite, non-negative amount.")

    unit = (_text(values.get("rate_unit")) or "").lower()
    unit = unit.replace(" ", "_").replace("-", "_")
    aliases = {
        "marla": "per_marla",
        "kanal": "per_kanal",
        "acre": "per_acre",
        "kila": "per_kila",
        "square_foot": "per_square_foot",
        "square_yard": "per_square_yard",
        "square_meter": "per_square_meter",
        "sqft": "per_square_foot",
        "square_feet": "per_square_foot",
        "per_sq_ft": "per_square_foot",
        "sqyd": "per_square_yard",
        "square_yards": "per_square_yard",
        "per_sq_yd": "per_square_yard",
        "sqm": "per_square_meter",
        "square_meters": "per_square_meter",
    }
    unit = aliases.get(unit, unit)
    if unit not in DC_RATE_UNITS:
        raise DCRateServiceError(
            "Choose a supported rate unit: " + ", ".join(DC_RATE_UNITS) + "."
        )

    raw_year = values.get("published_year")
    try:
        year = int(str(raw_year).strip()) if raw_year not in (None, "") else None
    except (TypeError, ValueError):
        raise DCRateServiceError("Published year must be a four-digit year.") from None
    if year is not None and not 1900 <= year <= 2200:
        raise DCRateServiceError("Published year must be between 1900 and 2200.")
    start = _date(values.get("effective_from"), "Effective from")
    if year is not None:
        start = date(year, 1, 1)
    if start is None:
        # effective_from is retained as a required legacy storage column. A
        # missing schedule year remains NULL so later valuation can flag it.
        start = date.today()
    elif year is None:
        year = start.year
    end = _date(values.get("effective_to"), "Effective to")
    source_date = _date(values.get("source_date"), "Source date")
    if end and end < start:
        raise DCRateServiceError("Effective to cannot be before Effective from.")

    status = (_text(values.get("verification_status")) or "unverified").lower()
    schedule_type = (_text(values.get("schedule_type")) or "dc").lower()
    if schedule_type not in {"dc", "fbr"}:
        raise DCRateServiceError("Schedule must be DC or FBR.")
    if status not in {"unverified", "verified"}:
        raise DCRateServiceError("Verification status must be Unverified or Verified.")
    source = _text(values.get("source_reference"))
    verified_by = _text(values.get("verified_by"))
    if status == "verified" and (not source or not verified_by):
        raise DCRateServiceError(
            "A verified rate requires a source reference and the verifier's name."
        )

    return {
        "schedule_type": schedule_type,
        "property_type": property_type,
        "classification": _text(values.get("classification")),
        "floor": _text(values.get("floor")),
        "rate_value": amount,
        "rate_unit": unit,
        "effective_from": start,
        "published_year": year,
        "effective_to": end,
        "source_reference": source,
        "source_date": source_date,
        "verification_status": status,
        "verified_at": datetime.now(timezone.utc) if status == "verified" else None,
        "verified_by": verified_by if status == "verified" else None,
    }


def _rate_identity(rate: dict[str, Any], location_key: tuple[str, ...]) -> tuple[Any, ...]:
    return (
        rate.get("schedule_type", "dc"),
        location_key,
        rate["property_type"].casefold(),
        (rate.get("classification") or "").casefold(),
        (rate.get("floor") or "").casefold(),
        rate["rate_unit"],
        rate["effective_from"],
        rate["effective_to"],
    )


def _periods_overlap(
    first_start: date,
    first_end: date | None,
    second_start: date,
    second_end: date | None,
) -> bool:
    """Treat an absent end date as an open-ended effective period."""
    return (first_end is None or second_start <= first_end) and (
        second_end is None or first_start <= second_end
    )


def _record_dict(rate: DCValuationRate) -> dict[str, Any]:
    location = rate.location
    return {
        "id": rate.id,
        "location_id": rate.location_id,
        "location_name": location.name if location else "Unassigned",
        "province": location.province if location else None,
        "district": location.district if location else None,
        "tehsil": location.tehsil if location else None,
        "land_type": location.land_type if location else None,
        "committee_town": location.committee_town if location else None,
        "revenue_circle": location.revenue_circle if location else None,
        "property_area": location.property_area if location else None,
        "qanoongoi": location.qanoongoi if location else None,
        "patwar_circle": location.patwar_circle if location else None,
        "mauza": location.mauza if location else None,
        "hadbast_number": location.hadbast_number if location else None,
        "square_no": location.square_no if location else None,
        "kila_no": location.kila_no if location else None,
        "khasra_no": location.khasra_no if location else None,
        "locality": location.locality if location else None,
        "property_type": rate.property_type,
        "schedule_type": rate.schedule_type,
        "classification": rate.classification,
        "floor": rate.floor,
        "rate_value": rate.rate_value,
        "rate_unit": rate.rate_unit,
        "effective_from": rate.effective_from,
        "published_year": rate.published_year,
        "effective_to": rate.effective_to,
        "source_reference": rate.source_reference,
        "source_date": rate.source_date,
        "verification_status": rate.verification_status,
        "verified_at": rate.verified_at,
        "verified_by": rate.verified_by,
        "created_at": rate.created_at,
        "updated_at": rate.updated_at,
    }


SCHEDULE_EXPORT_COLUMNS = (
    ("schedule_type", "Schedule"),
    ("district", "District"),
    ("tehsil", "Tehsil"),
    ("land_type", "Rural or urban"),
    ("committee_town", "Town / city"),
    ("revenue_circle", "Revenue circle"),
    ("qanoongoi", "Qanoongoi"),
    ("patwar_circle", "Patwar circle"),
    ("mauza", "Mauza"),
    ("square_no", "Square no"),
    ("kila_no", "Kila no"),
    ("khasra_no", "Khasra no"),
    ("floor", "Floor"),
    ("property_area", "Property area"),
    ("location_name", "Location"),
    ("classification", "Land classification"),
    ("rate_value", "Schedule rate"),
    ("rate_unit", "Unit"),
    ("published_year", "Published year"),
    ("source_reference", "Schedule reference"),
)


def export_schedule_records(
    records: list[dict[str, Any]], file_format: str, fields: list[str] | None = None
) -> bytes:
    """Serialize the selected schedule dimensions to CSV, XLSX, or ODS bytes."""
    columns_by_key = dict(SCHEDULE_EXPORT_COLUMNS)
    selected = fields or [key for key, _label in SCHEDULE_EXPORT_COLUMNS]
    unknown = set(selected) - columns_by_key.keys()
    if unknown:
        raise DCRateServiceError("Unsupported export field(s): " + ", ".join(sorted(unknown)))
    labels = [columns_by_key[key] for key in selected]

    def cell(record: dict[str, Any], key: str) -> str:
        value = record.get(key)
        if value is None:
            return ""
        return value.isoformat() if hasattr(value, "isoformat") else str(value)

    rows = [[cell(record, key) for key in selected] for record in records]
    output = BytesIO()
    if file_format == "csv":
        text_output = StringIO(newline="")
        writer = csv.writer(text_output)
        writer.writerow(labels)
        writer.writerows(rows)
        return text_output.getvalue().encode("utf-8-sig")
    if file_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Valuation Schedules"
        sheet.append(labels)
        for row in rows:
            sheet.append(row)
        workbook.save(output)
        workbook.close()
        return output.getvalue()
    if file_format == "ods":
        try:
            from odf.opendocument import OpenDocumentSpreadsheet
            from odf.table import Table, TableCell, TableRow
            from odf.text import P
            document = OpenDocumentSpreadsheet()
            sheet = Table(name="Valuation Schedules")
            for values in [labels, *rows]:
                row = TableRow()
                for value in values:
                    cell_control = TableCell()
                    if value:
                        cell_control.addElement(P(text=value))
                    row.addElement(cell_control)
                sheet.addElement(row)
            document.spreadsheet.addElement(sheet)
            document.write(output)
            return output.getvalue()
        except Exception as exc:
            raise DCRateServiceError(f"Could not create ODS export: {exc}") from exc
    raise DCRateServiceError("Choose CSV, XLSX, or ODS for export.")


def list_locations() -> list[dict[str, Any]]:
    with SessionLocal() as session:
        locations = session.scalars(
            select(PropertyLocation).order_by(
                PropertyLocation.district,
                PropertyLocation.tehsil,
                PropertyLocation.name,
                PropertyLocation.id,
            )
        ).all()
        return [
            {
                "id": location.id,
                **{field: getattr(location, field) for field in LOCATION_FIELDS},
            }
            for location in locations
        ]


def create_location(**values: Any) -> tuple[int, bool]:
    """Return (location ID, created); exact duplicates reuse the location."""
    normalized = _location_data(values)
    with SessionLocal.begin() as session:
        for existing in session.scalars(select(PropertyLocation)).all():
            existing_data = {field: getattr(existing, field) for field in LOCATION_FIELDS}
            if _location_key(existing_data) == _location_key(normalized):
                return existing.id, False
        location = PropertyLocation(**normalized)
        session.add(location)
        session.flush()
        return location.id, True


def _duplicate_rate(
    session,
    rate_data: dict[str, Any],
    location_id: int,
    *,
    exclude_id: int | None = None,
) -> DCValuationRate | None:
    statement = select(DCValuationRate).where(
        DCValuationRate.location_id == location_id,
        DCValuationRate.rate_unit == rate_data["rate_unit"],
    )
    if exclude_id is not None:
        statement = statement.where(DCValuationRate.id != exclude_id)
    for existing in session.scalars(statement).all():
        if (
            existing.schedule_type == rate_data.get("schedule_type", "dc")
            and existing.property_type.casefold() == rate_data["property_type"].casefold()
            and (existing.classification or "").casefold()
            == (rate_data.get("classification") or "").casefold()
            and (existing.floor or "").casefold()
            == (rate_data.get("floor") or "").casefold()
            and _periods_overlap(
                existing.effective_from,
                existing.effective_to,
                rate_data["effective_from"],
                rate_data["effective_to"],
            )
        ):
            return existing
    return None


def create_dc_rate(*, location_id: int, **values: Any) -> int:
    rate_data = _rate_data(values)
    with SessionLocal.begin() as session:
        if session.get(PropertyLocation, location_id) is None:
            raise DCRateServiceError("Select a valid property location.")
        duplicate = _duplicate_rate(session, rate_data, location_id)
        if duplicate:
            raise DCRateServiceError(
                f"This rate conflicts with existing record #{duplicate.id}; "
                "existing rates are never overwritten."
            )
        record = DCValuationRate(location_id=location_id, **rate_data)
        session.add(record)
        session.flush()
        return record.id


def update_dc_rate(*, rate_id: int, location_id: int, **values: Any) -> None:
    rate_data = _rate_data(values)
    with SessionLocal.begin() as session:
        record = session.get(DCValuationRate, rate_id)
        if record is None:
            raise DCRateServiceError("DC rate record not found.")
        if session.get(PropertyLocation, location_id) is None:
            raise DCRateServiceError("Select a valid property location.")
        duplicate = _duplicate_rate(
            session, rate_data, location_id, exclude_id=rate_id
        )
        if duplicate:
            raise DCRateServiceError(
                f"This rate conflicts with existing record #{duplicate.id}; "
                "existing rates are never overwritten."
            )
        record.location_id = location_id
        for field, value in rate_data.items():
            setattr(record, field, value)


def get_dc_rate(rate_id: int) -> dict[str, Any] | None:
    with SessionLocal() as session:
        record = session.scalar(
            select(DCValuationRate)
            .options(joinedload(DCValuationRate.location))
            .where(DCValuationRate.id == rate_id)
        )
        return _record_dict(record) if record else None


def list_dc_rates(
    *,
    search: str = "",
    schedule_type: str | None = None,
    location_id: int | None = None,
    rate_unit: str | None = None,
    verification_status: str | None = None,
    effective_on: date | None = None,
    published_year: int | None = None,
) -> list[dict[str, Any]]:
    statement = select(DCValuationRate).options(joinedload(DCValuationRate.location))
    if schedule_type in {"dc", "fbr"}:
        statement = statement.where(DCValuationRate.schedule_type == schedule_type)
    if location_id is not None:
        statement = statement.where(DCValuationRate.location_id == location_id)
    if rate_unit:
        statement = statement.where(DCValuationRate.rate_unit == rate_unit)
    if verification_status in {"unverified", "verified"}:
        statement = statement.where(
            DCValuationRate.verification_status == verification_status
        )
    if effective_on:
        statement = statement.where(
            DCValuationRate.effective_from <= effective_on,
            or_(DCValuationRate.effective_to.is_(None), DCValuationRate.effective_to >= effective_on),
        )
    if published_year is not None:
        statement = statement.where(DCValuationRate.published_year == published_year)
    search = search.strip()
    if search:
        pattern = f"%{search}%"
        statement = statement.join(
            PropertyLocation, DCValuationRate.location_id == PropertyLocation.id,
            isouter=True,
        ).where(
            or_(
                DCValuationRate.property_type.ilike(pattern),
                DCValuationRate.classification.ilike(pattern),
                DCValuationRate.floor.ilike(pattern),
                DCValuationRate.schedule_type.ilike(pattern),
                DCValuationRate.source_reference.ilike(pattern),
                PropertyLocation.name.ilike(pattern),
                PropertyLocation.district.ilike(pattern),
                PropertyLocation.tehsil.ilike(pattern),
                PropertyLocation.land_type.ilike(pattern),
                PropertyLocation.committee_town.ilike(pattern),
                PropertyLocation.revenue_circle.ilike(pattern),
                PropertyLocation.property_area.ilike(pattern),
                PropertyLocation.qanoongoi.ilike(pattern),
                PropertyLocation.patwar_circle.ilike(pattern),
                PropertyLocation.mauza.ilike(pattern),
                PropertyLocation.hadbast_number.ilike(pattern),
                PropertyLocation.square_no.ilike(pattern),
                PropertyLocation.kila_no.ilike(pattern),
                PropertyLocation.khasra_no.ilike(pattern),
            )
        )
    statement = statement.order_by(
        DCValuationRate.effective_from.desc(), DCValuationRate.id.desc()
    )
    with SessionLocal() as session:
        return [_record_dict(row) for row in session.scalars(statement).unique().all()]


def _header_key(value: Any) -> str:
    return (str(value or "").strip().casefold().replace("/", "_")
            .replace(" ", "_").replace("-", "_").replace("__", "_"))


def read_schedule_spreadsheet(file_bytes: bytes, filename: str) -> tuple[list[str], list[list[Any]]]:
    """Read tabular values from CSV, XLSX or ODS without changing source data."""
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == "csv":
        try:
            lines = file_bytes.decode("utf-8-sig").splitlines()
            rows = list(csv.reader(lines))
        except (UnicodeDecodeError, csv.Error) as exc:
            raise DCRateServiceError(f"Could not read CSV file: {exc}") from exc
    elif extension == "xlsx":
        try:
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
            rows = list(workbook.active.iter_rows(values_only=True))
            workbook.close()
        except Exception as exc:
            raise DCRateServiceError(f"Could not open Excel workbook: {exc}") from exc
    elif extension == "ods":
        try:
            from odf.opendocument import load as load_ods
            from odf.table import Table, TableCell, TableRow
            from odf.text import P
            from odf import teletype
            document = load_ods(BytesIO(file_bytes))
            sheets = document.spreadsheet.getElementsByType(Table)
            if not sheets:
                raise DCRateServiceError("The ODS file has no worksheet.")
            rows = []
            for row in sheets[0].getElementsByType(TableRow):
                values = []
                for cell in row.getElementsByType(TableCell):
                    repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
                    raw = cell.getAttribute("datevalue") or cell.getAttribute("value")
                    if raw is None:
                        raw = teletype.extractText(cell)
                    values.extend([raw] * min(repeat, 256))
                rows.append(values)
        except DCRateServiceError:
            raise
        except Exception as exc:
            raise DCRateServiceError(f"Could not open ODS workbook: {exc}") from exc
    else:
        raise DCRateServiceError("Choose a CSV, XLSX, or ODS spreadsheet.")
    rows = [list(row) for row in rows if any(value is not None and str(value).strip() for value in row)]
    if not rows:
        raise DCRateServiceError("The selected spreadsheet is empty.")
    headers = [str(value or "").strip() for value in rows[0]]
    return headers, rows[1:]


def suggested_column_mapping(headers: list[str]) -> dict[str, int]:
    """Suggest mappings from familiar schedule headers; user confirms them."""
    mapping = {}
    for index, header in enumerate(headers):
        target = RATE_HEADERS.get(_header_key(header))
        if _header_key(header) == "location":
            target = "locality"
        if target and target not in mapping:
            mapping[target] = index
    # The current official schedule exports usually expose Mauza and Location,
    # but not a separate display-name column for the database location record.
    if "location_name" not in mapping:
        if "mauza" in mapping:
            mapping["location_name"] = mapping["mauza"]
        elif "locality" in mapping:
            mapping["location_name"] = mapping["locality"]
    return mapping


def preview_mapped_schedule(
    headers: list[str], rows: list[list[Any]], mapping: dict[str, int], schedule_type: str
) -> dict[str, Any]:
    """Normalize user-mapped columns, then reuse the validated import pipeline."""
    if schedule_type not in {"dc", "fbr"}:
        raise DCRateServiceError("Select either the DC or FBR schedule.")
    mapping = dict(mapping)
    # In many published schedules the only class column is Land Classification.
    # It also supplies the broad property type when a separate column is absent.
    if "property_type" not in mapping and "classification" in mapping:
        mapping["property_type"] = mapping["classification"]
    if "location_name" not in mapping:
        if "mauza" in mapping:
            mapping["location_name"] = mapping["mauza"]
        elif "locality" in mapping:
            mapping["location_name"] = mapping["locality"]
    missing = REQUIRED_RATE_COLUMNS - mapping.keys()
    if "classification" in mapping:
        missing.discard("property_type")
    if missing:
        raise DCRateServiceError("Map required columns: " + ", ".join(sorted(missing)))
    output = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    fields = list(mapping) + ["schedule_type"]
    sheet.append(fields)
    for row in rows:
        sheet.append([
            *[row[index] if index < len(row) else None for index in mapping.values()],
            schedule_type,
        ])
    workbook.save(output)
    workbook.close()
    preview = preview_dc_rate_workbook(output.getvalue())
    return preview


def preview_dc_rate_workbook(file_bytes: bytes) -> dict[str, Any]:
    """Validate an XLSX sheet and report invalid/duplicate rows without writes."""
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise DCRateServiceError(f"Could not open this Excel workbook: {exc}") from exc
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration:
            raise DCRateServiceError("The selected worksheet is empty.") from None

        columns: dict[str, int] = {}
        duplicate_headers = []
        for index, value in enumerate(header_row):
            raw_header = _header_key(value)
            if not raw_header:
                continue
            target = RATE_HEADERS.get(raw_header)
            if target:
                if target in columns:
                    duplicate_headers.append(target)
                columns[target] = index
        if duplicate_headers:
            raise DCRateServiceError(
                "The worksheet contains duplicate columns: "
                + ", ".join(sorted(set(duplicate_headers)))
            )
        missing = REQUIRED_RATE_COLUMNS - columns.keys()
        if "classification" in columns:
            missing.discard("property_type")
        if missing:
            raise DCRateServiceError(
                "Missing required columns: " + ", ".join(sorted(missing))
            )

        candidates: list[dict[str, Any]] = []
        issues: list[dict[str, Any]] = []
        seen: dict[tuple[Any, ...], list[tuple[date, date | None, int]]] = {}
        rows_read = 0
        for excel_row, values in enumerate(iterator, start=2):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            rows_read += 1
            raw = {
                field: values[index] if index < len(values) else None
                for field, index in columns.items()
            }
            try:
                location = _location_data({
                    "name": raw.get("location_name"),
                    "province": raw.get("province"),
                    "district": raw.get("district"),
                    "tehsil": raw.get("tehsil"),
                    "land_type": raw.get("land_type"),
                    "committee_town": raw.get("committee_town"),
                    "revenue_circle": raw.get("revenue_circle"),
                    "property_area": raw.get("property_area"),
                    "qanoongoi": raw.get("qanoongoi"),
                    "patwar_circle": raw.get("patwar_circle"),
                    "mauza": raw.get("mauza"),
                    "hadbast_number": raw.get("hadbast_number"),
                    "square_no": raw.get("square_no"),
                    "kila_no": raw.get("kila_no"),
                    "khasra_no": raw.get("khasra_no"),
                    "locality": raw.get("locality"),
                })
                rate = _rate_data(raw)
                identity = _rate_identity(rate, _location_key(location))
                scope = identity[:6]
                conflict_row = next(
                    (
                        prior_row
                        for prior_start, prior_end, prior_row in seen.get(scope, [])
                        if _periods_overlap(
                            prior_start,
                            prior_end,
                            rate["effective_from"],
                            rate["effective_to"],
                        )
                    ),
                    None,
                )
                if conflict_row is not None:
                    issues.append({
                        "row": excel_row,
                        "kind": "duplicate in workbook",
                        "message": f"Conflicting effective period with row {conflict_row} in this workbook.",
                    })
                    continue
                seen.setdefault(scope, []).append(
                    (rate["effective_from"], rate["effective_to"], excel_row)
                )
                candidates.append({"row": excel_row, "location": location, **rate})
            except DCRateServiceError as exc:
                issues.append({"row": excel_row, "kind": "invalid", "message": str(exc)})

        with SessionLocal() as session:
            existing_keys: list[tuple[tuple[Any, ...], int]] = []
            locations = session.scalars(select(PropertyLocation)).all()
            location_by_id = {location.id: location for location in locations}
            rates = session.scalars(select(DCValuationRate)).all()
            for rate in rates:
                location = location_by_id.get(rate.location_id)
                if location is not None:
                    location_values = {
                        field: getattr(location, field) for field in LOCATION_FIELDS
                    }
                    existing_keys.append(
                        (
                            _rate_identity(
                                {
                                    "schedule_type": rate.schedule_type,
                                    "property_type": rate.property_type,
                                    "classification": rate.classification,
                                    "floor": rate.floor,
                                    "rate_unit": rate.rate_unit,
                                    "effective_from": rate.effective_from,
                                    "effective_to": rate.effective_to,
                                },
                                _location_key(location_values),
                            ),
                            rate.id,
                        )
                    )

        new_candidates = []
        for candidate in candidates:
            identity = _rate_identity(candidate, _location_key(candidate["location"]))
            conflict_id = next(
                (
                    rate_id
                    for existing_key, rate_id in existing_keys
                    if existing_key[:6] == identity[:6]
                    and _periods_overlap(
                        existing_key[6],
                        existing_key[7],
                        identity[6],
                        identity[7],
                    )
                ),
                None,
            )
            if conflict_id is not None:
                issues.append({
                    "row": candidate["row"],
                    "kind": "duplicate in database",
                    "message": (
                        f"Conflicts with existing record #{conflict_id} for the "
                        "same location, property class, unit, and effective period."
                    ),
                })
                continue
            new_candidates.append(candidate)
        return {
            "valid_records": new_candidates,
            "issues": sorted(issues, key=lambda issue: issue["row"]),
            "rows_read": rows_read,
        }
    finally:
        workbook.close()


def import_dc_rate_records(records: list[dict[str, Any]]) -> dict[str, Any]:
    """Insert valid preview rows, skipping any duplicate found at commit time."""
    imported_ids: list[int] = []
    issues: list[dict[str, Any]] = []
    with SessionLocal.begin() as session:
        for candidate in records:
            row = candidate["row"]
            location_values = _location_data(candidate["location"])
            rate_values = _rate_data(candidate)
            location = next(
                (
                    existing
                    for existing in session.scalars(select(PropertyLocation)).all()
                    if _location_key(
                        {field: getattr(existing, field) for field in LOCATION_FIELDS}
                    ) == _location_key(location_values)
                ),
                None,
            )
            if location is None:
                location = PropertyLocation(**location_values)
                session.add(location)
                session.flush()
            duplicate = _duplicate_rate(session, rate_values, location.id)
            if duplicate:
                issues.append({
                    "row": row,
                    "kind": "duplicate at import time",
                    "message": f"Conflicts with existing record #{duplicate.id}; skipped.",
                })
                continue
            record = DCValuationRate(location_id=location.id, **rate_values)
            session.add(record)
            session.flush()
            imported_ids.append(record.id)
    return {"imported_count": len(imported_ids), "issues": issues}
