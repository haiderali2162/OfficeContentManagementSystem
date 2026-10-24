
"""
Land Area Export Service
Version: 1.0.0

Exports saved area-calculation history to CSV, XLSX, ODS and PDF.
"""

import csv
import json
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from odf.opendocument import OpenDocumentSpreadsheet
from odf.style import Style, TableColumnProperties, TableCellProperties
from odf.table import Table, TableColumn, TableRow, TableCell
from odf.text import P

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table as PDFTable,
    TableStyle,
    PageBreak,
    KeepTogether,
)

from core.config import DATA_DIR
from services.area_history_service import list_calculations


EXPORT_DIR = DATA_DIR / "exports"

EXPORT_COLUMNS = [
    ("id", "Record ID"),
    ("created_at", "Date and Time"),
    ("operation", "Operation"),
    ("input_format", "Input Format"),
    ("inputs", "Original Inputs"),
    ("settings", "Measurement Settings"),
    ("result", "Result"),
    ("details", "Calculation Working"),
]


def _timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def _get_records() -> list[dict]:
    """Fetch saved history for export."""
    return list_calculations(limit=100000)


def _json_text(value) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _row_values(record: dict) -> list[str]:
    """Prepare a consistent row for every export format."""
    return [
        str(record.get("id", "")),
        str(record.get("created_at", "")),
        str(record.get("operation", "")),
        str(record.get("input_format", "")),
        _json_text(record.get("inputs", {})),
        _json_text(record.get("settings", {})),
        str(record.get("result", "")),
        str(record.get("details", "")),
    ]


def export_csv(records: list[dict] | None = None) -> Path:
    """Export calculation history as UTF-8 CSV."""
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    records = _get_records() if records is None else records

    path = EXPORT_DIR / f"area_history_{_timestamp()}.csv"

    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow([label for _, label in EXPORT_COLUMNS])

        for record in records:
            writer.writerow(_row_values(record))

    return path


def export_xlsx(records: list[dict] | None = None) -> Path:
    """Export history as a formatted Excel workbook."""
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    records = _get_records() if records is None else records

    path = EXPORT_DIR / f"area_history_{_timestamp()}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Calculation History"

    headers = [label for _, label in EXPORT_COLUMNS]
    sheet.append(headers)

    for record in records:
        sheet.append(_row_values(record))

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32

    widths = [12, 24, 20, 26, 42, 48, 55, 65]

    for column_number, width in enumerate(widths, start=1):
        sheet.column_dimensions[
            get_column_letter(column_number)
        ].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    workbook.save(path)
    return path


def export_ods(records: list[dict] | None = None) -> Path:
    """Export history as an OpenDocument spreadsheet."""
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    records = _get_records() if records is None else records

    path = EXPORT_DIR / f"area_history_{_timestamp()}.ods"

    document = OpenDocumentSpreadsheet()

    header_style = Style(
        name="HeaderCell",
        family="table-cell",
    )
    header_style.addElement(
        TableCellProperties(
            backgroundcolor="#1F4E78",
            padding="0.08in",
        )
    )
    document.automaticstyles.addElement(header_style)

    table = Table(name="Calculation History")

    for _ in EXPORT_COLUMNS:
        table.addElement(TableColumn())

    def add_cell(row, value, style=None):
        cell = TableCell(
            valuetype="string",
            stylename=style,
        )
        cell.addElement(P(text=str(value)))
        row.addElement(cell)

    header_row = TableRow()

    for _, label in EXPORT_COLUMNS:
        add_cell(header_row, label, "HeaderCell")

    table.addElement(header_row)

    for record in records:
        row = TableRow()

        for value in _row_values(record):
            add_cell(row, value)

        table.addElement(row)

    document.spreadsheet.addElement(table)
    document.save(str(path))

    return path


def export_pdf(records: list[dict] | None = None) -> Path:
    """Export history as a print-ready PDF report."""
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    records = _get_records() if records is None else records

    path = EXPORT_DIR / f"area_history_{_timestamp()}.pdf"

    page_size = landscape(A4)

    document = SimpleDocTemplate(
        str(path),
        pagesize=page_size,
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=18 * mm,
        bottomMargin=17 * mm,
        title="Land Area Calculation History",
        author="Land Office Manager",
    )

    base_styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ReportTitle",
        parent=base_styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=17,
        leading=21,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1F4E78"),
        spaceAfter=5 * mm,
    )

    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=base_styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#555555"),
        spaceAfter=4 * mm,
    )

    cell_style = ParagraphStyle(
        "ReportCell",
        parent=base_styles["Normal"],
        fontName="Helvetica",
        fontSize=7,
        leading=9,
        alignment=TA_LEFT,
        wordWrap="LTR",
    )

    heading_style = ParagraphStyle(
        "RecordHeading",
        parent=base_styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#1F4E78"),
        spaceBefore=3 * mm,
        spaceAfter=2 * mm,
    )

    body_style = ParagraphStyle(
        "RecordBody",
        parent=base_styles["BodyText"],
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        spaceAfter=2 * mm,
        wordWrap="LTR",
    )

    story = [
        Paragraph("Land Area Calculation History", title_style),
        Paragraph(
            f"Generated: {escape(datetime.now().strftime('%Y-%m-%d %H:%M:%S'))}",
            subtitle_style,
        ),
        Paragraph(
            f"Total saved calculations in this report: {len(records)}",
            subtitle_style,
        ),
        Spacer(1, 3 * mm),
    ]

    if not records:
        story.append(
            Paragraph("No calculation history is available.", body_style)
        )
    else:
        # Summary table: short fields only, to keep the report readable.
        summary_headers = [
            "ID",
            "Date and Time",
            "Operation",
            "Input Format",
            "Result",
        ]

        summary_data = [
            [Paragraph(escape(value), cell_style) for value in summary_headers]
        ]

        for record in records:
            result_summary = str(record.get("result", ""))

            # Keep the summary compact; full results appear in the details below.
            result_summary = result_summary.splitlines()[0][:160]

            values = [
                str(record.get("id", "")),
                str(record.get("created_at", "")),
                str(record.get("operation", "")),
                str(record.get("input_format", "")),
                result_summary,
            ]

            summary_data.append([
                Paragraph(escape(value), cell_style)
                for value in values
            ])

        summary_table = PDFTable(
            summary_data,
            colWidths=[
                12 * mm,
                34 * mm,
                28 * mm,
                38 * mm,
                150 * mm,
            ],
            repeatRows=1,
            hAlign="LEFT",
        )

        summary_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7C9D6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
                    colors.white,
                    colors.HexColor("#F2F6F9"),
                ]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ])
        )

        story.append(summary_table)
        story.append(PageBreak())
        story.append(
            Paragraph("Detailed Calculation Records", title_style)
        )

        for index, record in enumerate(records):
            record_id = escape(str(record.get("id", "")))

            story.append(
                Paragraph(
                    f"Calculation #{record_id} — "
                    f"{escape(str(record.get('operation', '')))}",
                    heading_style,
                )
            )

            metadata = [
                ("Date and time", record.get("created_at", "")),
                ("Input format", record.get("input_format", "")),
                ("Original inputs", _json_text(record.get("inputs", {}))),
                ("Measurement settings", _json_text(record.get("settings", {}))),
            ]

            metadata_data = []

            for label, value in metadata:
                metadata_data.append([
                    Paragraph(f"<b>{escape(label)}</b>", cell_style),
                    Paragraph(escape(str(value)), cell_style),
                ])

            metadata_table = PDFTable(
                metadata_data,
                colWidths=[42 * mm, 220 * mm],
                hAlign="LEFT",
            )

            metadata_table.setStyle(
                TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CDD7DF")),
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F6F9")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ])
            )

            story.append(metadata_table)
            story.append(Spacer(1, 2 * mm))
            story.append(Paragraph("Result", heading_style))

            result_lines = str(record.get("result", "")).splitlines()

            for line in result_lines:
                story.append(
                    Paragraph(escape(line) or " ", body_style)
                )

            story.append(Paragraph("Calculation working", heading_style))

            detail_lines = str(record.get("details", "")).splitlines()

            for line in detail_lines:
                story.append(
                    Paragraph(escape(line) or " ", body_style)
                )

            if index < len(records) - 1:
                story.append(Spacer(1, 4 * mm))

    def draw_page_number(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#B7C9D6"))
        canvas.line(
            doc.leftMargin,
            12 * mm,
            page_size[0] - doc.rightMargin,
            12 * mm,
        )
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(
            doc.leftMargin,
            7 * mm,
            "Land Office Manager — Calculation History",
        )
        canvas.drawRightString(
            page_size[0] - doc.rightMargin,
            7 * mm,
            f"Page {doc.page}",
        )
        canvas.restoreState()

    document.build(
        story,
        onFirstPage=draw_page_number,
        onLaterPages=draw_page_number,
    )

    return path


def export_area_history(file_format: str) -> Path:
    """Export all available saved calculations in the requested format."""
    normalized = file_format.strip().lower()
    records = _get_records()

    exporters = {
        "csv": export_csv,
        "xlsx": export_xlsx,
        "ods": export_ods,
        "pdf": export_pdf,
    }

    if normalized not in exporters:
        raise ValueError(
            "Unsupported format. Choose CSV, XLSX, ODS or PDF."
        )

    return exporters[normalized](records)
