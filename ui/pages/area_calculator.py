"""Land-area calculator with configurable local units and saved history."""

import traceback
import csv
import io
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import flet as ft
from openpyxl import load_workbook
from odf.opendocument import load as load_ods
from odf.table import Table, TableRow, TableCell
from odf.teletype import extractText

from core.area_engine import (
    AreaError,
    AreaSettings,
    LandArea,
    average_area,
    sum_areas,
    subtract_areas,
)
from services.area_export_service import export_area_history
from services.area_history_service import (
    get_calculation,
    list_calculations,
    save_calculation,
    clear_calculations,
)


DISPLAY_UNITS = ["Kanal-Marla-Foot", "Acres", "Kilas"]
OPERATIONS = [
    "Convert",
    "Add",
    "Subtract",
    "Multiply",
    "Divide",
    "Average",
    "Percentage Share",
]


def _format_decimal(value: Decimal, places: int = 6) -> str:
    """Format a decimal consistently without displaying unnecessary zeroes."""
    rounded = value.quantize(
        Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP
    )
    return format(rounded, "f").rstrip("0").rstrip(".") or "0"


def build_area_calculator_page(page: ft.Page) -> ft.Control:
    """Build calculator controls and connect them to the area engine/history."""
    marla_dropdown = ft.Dropdown(
        label="Square feet per marla",
        value="272",
        options=[
            ft.dropdown.Option("272"),
            ft.dropdown.Option("225"),
            ft.dropdown.Option("custom", "Custom"),
        ],
        expand=True,
    )
    custom_marla_field = ft.TextField(
        label="Custom square feet per marla",
        hint_text="Enter a positive whole number",
        keyboard_type=ft.KeyboardType.NUMBER,
        value="",
        expand=True,
        visible=False,
    )
    conversion_dropdown = ft.Dropdown(
        label="Show result as",
        value="Kanal-Marla-Foot",
        options=[ft.dropdown.Option(unit) for unit in DISPLAY_UNITS],
        expand=True,
    )
    operation_dropdown = ft.Dropdown(
        label="Operation",
        value="Convert",
        options=[ft.dropdown.Option(operation) for operation in OPERATIONS],
        expand=True,
    )

    # Area inputs use K, K-M, or K-M-F. Multiplication and division use B as
    # a numeric factor; area arithmetic writes a normalized K-M-F value to A.
    value_a = ft.TextField(
        label="A",
        hint_text="Examples: 3-10-120, 12 kilas, or 3 acres",
        expand=True,
        autofocus=True,
    )
    value_b = ft.TextField(
        label="B",
        hint_text="Area: K, K-M, or K-M-F",
        expand=True,
    )
    preview_a = ft.Text(
        "Enter an area as K, K-M, or K-M-F to preview its units.",
        size=12,
        selectable=True,
        color="#475569",
    )
    preview_b = ft.Text(
        "Enter an area as K, K-M, or K-M-F to preview its units.",
        size=12,
        selectable=True,
        color="#475569",
    )
    preview_a_card = ft.Container(
        content=ft.Column(
            controls=[ft.Text("A area preview", weight=ft.FontWeight.BOLD), preview_a],
            spacing=6,
        ),
        padding=12,
        bgcolor="#F8FAFC",
        border=ft.Border.all(1, "#E2E8F0"),
        border_radius=10,
        expand=True,
    )
    preview_a_card.col = {"sm": 6, "md": 6}
    preview_b_card = ft.Container(
        content=ft.Column(
            controls=[ft.Text("B area preview", weight=ft.FontWeight.BOLD), preview_b],
            spacing=6,
        ),
        padding=12,
        bgcolor="#F8FAFC",
        border=ft.Border.all(1, "#E2E8F0"),
        border_radius=10,
        expand=True,
    )
    preview_b_card.col = {"sm": 6, "md": 6}
    calculate_button = ft.Button("Calculate")
    clear_button = ft.OutlinedButton("Clear")
    status_text = ft.Text("", color="#B91C1C", size=12)
    result_text = ft.Text(
        "Enter K, K-M, or K-M-F in A and choose an operation.",
        selectable=True,
    )
    details_text = ft.Text("", selectable=True, size=12, color="#475569")
    history_column = ft.Column(spacing=8, scroll=ft.ScrollMode.AUTO)
    file_picker = ft.FilePicker()
    # FilePicker is a Flet page service; attach once so native and web builds
    # can use the platform's file dialogs for both imports and exports.
    if not any(isinstance(service, ft.FilePicker) for service in page.services):
        page.services.append(file_picker)
    else:
        file_picker = next(
            service for service in page.services if isinstance(service, ft.FilePicker)
        )

    def make_settings() -> AreaSettings:
        """Create the engine settings from the marla control and convention."""
        selected = marla_dropdown.value or "272"
        if selected == "custom":
            raw_custom_size = (custom_marla_field.value or "").strip()
            try:
                custom_size = int(raw_custom_size)
            except (ValueError, TypeError):
                raise AreaError(
                    "Enter a whole-number square-foot size for one marla."
                ) from None
            if str(custom_size) != raw_custom_size or custom_size <= 0:
                raise AreaError("Custom square feet per marla must be positive.")
            marla_size = custom_size
        else:
            marla_size = int(selected)

        # The local convention is 20 marlas per kanal, 8 kanals per acre,
        # and one kila = 8 kanals + 18 marlas.
        return AreaSettings(
            square_feet_per_marla=marla_size,
            marla_per_kanal=20,
            sarsahi_per_marla=9,
            kanals_per_acre=8,
            kila_kanal=8,
            kila_marla=18,
            kila_feet=0,
            feet_per_karam=Decimal("5.5"),
        )

    def parse_area(raw_value: str, label: str, settings: AreaSettings) -> LandArea:
        """Accept K-M-F input or an amount explicitly labeled acres or kilas."""
        raw = (raw_value or "").strip()
        unit_value = re.fullmatch(
            r"(\d+(?:\.\d+)?)\s*(acres?|kilas?)", raw, flags=re.IGNORECASE
        )
        if unit_value:
            amount, unit = unit_value.groups()
            if unit.lower().startswith("acre"):
                return LandArea.from_acres(amount, settings)
            return LandArea.from_kila(amount, settings)
        parts = raw.split("-")
        if not raw or len(parts) not in {1, 2, 3}:
            raise AreaError(
                f"{label} must use K, K-M, or K-M-F format, such as 2-5-100."
            )
        try:
            values = [int(part.strip()) for part in parts]
        except (ValueError, TypeError):
            raise AreaError(
                f"{label} must use whole-number K, M, and F components."
            ) from None
        values.extend([0] * (3 - len(values)))
        return LandArea.from_kmf(*values, settings=settings)

    def parse_number(raw_value: str, label: str) -> Decimal:
        """Validate a scalar B operand as a finite, non-negative decimal."""
        try:
            value = Decimal((raw_value or "").strip())
        except (InvalidOperation, ValueError, TypeError):
            raise AreaError(f"Enter a valid number for {label}.") from None
        if not value.is_finite() or value < 0:
            raise AreaError(f"{label} must be a finite, non-negative value.")
        return value

    def parse_share(raw_value: str) -> Decimal:
        """Convert a fraction (1/5) or percent (20) to percentage points."""
        raw = (raw_value or "").strip()
        fraction = re.fullmatch(r"(\d+)\s*/\s*(\d+)", raw)
        if fraction:
            numerator, denominator = (int(part) for part in fraction.groups())
            if denominator == 0:
                raise AreaError("Share denominator must be greater than zero.")
            if numerator > denominator:
                raise AreaError("The share cannot exceed the total area.")
            return Decimal(numerator) * Decimal(100) / Decimal(denominator)
        if raw.endswith("%"):
            raw = raw[:-1].strip()
        percentage = parse_number(raw, "B percentage")
        if percentage > 100:
            raise AreaError("Percentage must be between 0 and 100.")
        return percentage

    def format_area_preview(raw_value: str, settings: AreaSettings) -> str:
        """Build live unit details for one syntactically complete K-M-F value."""
        if not (raw_value or "").strip():
            return "Type K, K-M, or K-M-F to see its area conversions."
        try:
            area = parse_area(raw_value, "Area", settings)
        except AreaError:
            # While users are midway through typing a hyphenated value, keep
            # the preview calm and descriptive instead of showing an error.
            return "Complete the value as K, K-M, or K-M-F to see conversions."

        total_marlas = area.square_feet / Decimal(settings.square_feet_per_marla)
        total_sarsahi = total_marlas * Decimal(settings.sarsahi_per_marla)
        return (
            f"Acres: {_format_decimal(area.to_acres(settings))} | "
            f"Kilas: {_format_decimal(area.to_kila(settings))}\n"
            f"Square feet: {_format_decimal(area.square_feet, places=2)} | "
            f"Marlas: {_format_decimal(total_marlas)}\n"
            f"Sarsahi (9 per marla): {_format_decimal(total_sarsahi)} | "
            f"K-M-S: {area.format_sarsahi(settings)}\n"
            f"Square karams (5.5 ft per karam): "
            f"{_format_decimal(area.to_square_karam(settings))}"
        )

    def update_area_previews(e=None):
        """Refresh A/B conversions as either area field changes."""
        try:
            settings = make_settings()
            preview_a.value = format_area_preview(value_a.value or "", settings)
            if (operation_dropdown.value or "Convert") == "Convert":
                preview_b.value = "B is not used for the Convert operation."
            elif operation_dropdown.value in {"Multiply", "Divide", "Percentage Share"}:
                preview_b.value = (
                    "B is a numeric multiplier."
                    if operation_dropdown.value == "Multiply"
                    else "B is a numeric divisor."
                    if operation_dropdown.value == "Divide"
                    else "B can be a fraction such as 1/5 or a percentage."
                )
            else:
                preview_b.value = format_area_preview(
                    value_b.value or "", settings
                )
        except AreaError as exc:
            preview_a.value = str(exc)
            preview_b.value = str(exc)
        if e is not None:
            page.update()

    def value_from_area(area: LandArea, display_unit: str, settings: AreaSettings):
        """Render an area using the user's selected conversion destination."""
        if display_unit == "Kanal-Marla-Foot":
            return area.format(settings)
        if display_unit == "Acres":
            return area.to_acres(settings)
        if display_unit == "Kilas":
            return area.to_kila(settings)
        raise AreaError("Select K-M-F, Acres, or Kilas for the result display.")

    def format_result(area: LandArea, settings: AreaSettings, display_unit: str) -> str:
        """Show the normalized area plus clear acre and kila conversions."""
        marlas = area.square_feet / Decimal(settings.square_feet_per_marla)
        kanals = marlas / Decimal(settings.marla_per_kanal)
        displayed_value = value_from_area(area, display_unit, settings)
        formatted_value = (
            displayed_value
            if isinstance(displayed_value, str)
            else _format_decimal(displayed_value)
        )
        return (
            f"Result: {formatted_value} {display_unit}\n\n"
            f"K-M-F: {area.format(settings)}\n"
            f"Marlas: {_format_decimal(marlas)}\n"
            f"Kanals: {_format_decimal(kanals)}\n"
            f"Acres (8 kanals): {_format_decimal(area.to_acres(settings))}\n"
            f"Kilas (8 kanals + 18 marlas): "
            f"{_format_decimal(area.to_kila(settings))}\n"
            f"Square feet: {_format_decimal(area.square_feet, places=2)}"
        )

    def refresh_history():
        """Rebuild history cards from persisted area-calculation records."""
        history_column.controls.clear()
        try:
            records = list_calculations(100)
            if not records:
                history_column.controls.append(
                    ft.Text("No saved calculations yet.", color="#64748B")
                )
            else:
                for record in records:
                    title = f"#{record['id']} — {record['operation']}"
                    summary = (
                        f"{record['input_format']} · {record['result']}"
                    )

                    def open_record(e, record_id=record["id"]):
                        selected = get_calculation(record_id)
                        if selected is None:
                            status_text.value = "History record not found."
                        else:
                            result_text.value = selected["result"]
                            details_text.value = (
                                selected["details"]
                                + "\n\nSaved inputs:\n"
                                + str(selected["inputs"])
                                + "\n\nSaved settings:\n"
                                + str(selected["settings"])
                            )
                            status_text.value = (
                                f"Loaded calculation #{record_id}."
                            )
                        page.update()

                    history_column.controls.append(
                        ft.Container(
                            padding=12,
                            border=ft.Border.all(1, "#E2E8F0"),
                            border_radius=10,
                            bgcolor="#FFFFFF",
                            content=ft.Row(
                                controls=[
                                    ft.Column(
                                        controls=[
                                            ft.Text(
                                                title,
                                                weight=ft.FontWeight.BOLD,
                                            ),
                                            ft.Text(summary, selectable=True),
                                        ],
                                        expand=True,
                                        spacing=4,
                                    ),
                                    ft.TextButton(
                                        "Details",
                                        on_click=open_record,
                                    ),
                                ],
                                vertical_alignment=ft.CrossAxisAlignment.CENTER,
                            ),
                        )
                    )
        except Exception as exc:
            history_column.controls.append(
                ft.Text(f"Could not load history: {exc}", color="#B91C1C")
            )
        # Callers update the page after rebuilding this list. Avoid reading
        # history_column.page here because the initial render is not mounted.

    async def export_history(file_format: str):
        """Build an export, then let the user choose where to save it."""
        try:
            saved_path = export_area_history(file_format)
            destination = await file_picker.save_file(
                dialog_title=f"Save area history as {file_format.upper()}",
                file_name=Path(saved_path).name,
                allowed_extensions=[file_format],
                src_bytes=Path(saved_path).read_bytes(),
            )
            status_text.value = (
                f"{file_format.upper()} export saved to: {destination}"
                if destination
                else f"Export prepared in {saved_path} (save dialog cancelled)."
            )
        except Exception as exc:
            traceback.print_exc()
            status_text.value = f"Could not export {file_format.upper()}: {exc}"
        page.update()

    def make_export_handler(file_format: str):
        """Return a real async event handler so Flet awaits the Save dialog."""
        async def handle_export(e=None):
            await export_history(file_format)

        return handle_export

    export_buttons = ft.Row(
        controls=[
            ft.TextButton(
                f"Export {file_format.upper()}",
                on_click=make_export_handler(file_format),
            )
            for file_format in ("csv", "xlsx", "ods", "pdf")
        ],
        wrap=True,
        spacing=8,
    )

    async def import_area_list(e=None):
        """Read a one-area-value-per-row text or spreadsheet and total it."""
        try:
            selected_files = await file_picker.pick_files(
                dialog_title="Import area values",
                allowed_extensions=["txt", "csv", "xlsx", "ods"],
                allow_multiple=False,
                with_data=True,
            )
            if not selected_files:
                return
            selected = selected_files[0]
            if selected.bytes is not None:
                payload = selected.bytes
            elif selected.path:
                payload = Path(selected.path).read_bytes()
            else:
                raise ValueError("The selected file could not be read by the app.")
            suffix = Path(selected.name).suffix.lower()
            rows = []
            if suffix in {".txt", ".csv"}:
                decoded = payload.decode("utf-8-sig")
                if suffix == ".csv":
                    try:
                        dialect = csv.Sniffer().sniff(decoded[:4096])
                    except csv.Error:
                        dialect = csv.excel
                    rows = list(csv.reader(io.StringIO(decoded), dialect))
                else:
                    rows = [[line] for line in decoded.splitlines()]
            elif suffix == ".xlsx":
                workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
                sheet = workbook.active
                rows = [list(row) for row in sheet.iter_rows(values_only=True)]
                workbook.close()
            elif suffix == ".ods":
                document = load_ods(io.BytesIO(payload))
                sheet = document.spreadsheet.getElementsByType(Table)[0]
                for row in sheet.getElementsByType(TableRow):
                    cells = []
                    for cell in row.getElementsByType(TableCell):
                        repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
                        cells.extend([extractText(cell)] * min(repeat, 100))
                    rows.append(cells)
            else:
                raise ValueError("Choose a TXT, CSV, XLSX, or ODS file.")

            settings = make_settings()
            parsed = []
            skipped = 0
            errors = []
            header_words = {"area", "areas", "land area", "measurement", "k-m", "k-m-f", "kmf"}
            for row_number, row in enumerate(rows, start=1):
                # Lists commonly have one area per row. Taking the first
                # populated cell also tolerates spreadsheets with blank columns.
                candidate = next((cell for cell in row if cell is not None and str(cell).strip()), None)
                if candidate is None:
                    skipped += 1
                    continue
                raw = str(candidate).strip()
                if raw.lower() in header_words:
                    skipped += 1
                    continue
                # Spreadsheet numeric values represent whole kanal counts;
                # normalize Excel's 3.0 representation before the K-M-F parser.
                if isinstance(candidate, float) and candidate.is_integer():
                    raw = str(int(candidate))
                try:
                    parsed.append((raw, parse_area(raw, f"Row {row_number}", settings)))
                except AreaError as exc:
                    # Ignore ordinary title/header text, but report malformed
                    # area-like values so a typo cannot silently change totals.
                    if re.fullmatch(r"[\d\s.\-]+", raw):
                        errors.append(f"row {row_number}: {exc}")
                    else:
                        skipped += 1
            if errors:
                raise ValueError("Invalid area values: " + "; ".join(errors[:8]))
            if not parsed:
                raise ValueError("No K, K-M, or K-M-F area values were found.")
            total = sum_areas([area for _, area in parsed])
            answer = total.format(settings)
            value_a.value = answer
            value_b.value = ""
            operation_dropdown.value = "Convert"
            update_operation_controls()
            result_text.value = f"Imported total = {answer} K-M-F\n\n{format_result(total, settings, conversion_dropdown.value or 'Kanal-Marla-Foot')}"
            details_text.value = f"Imported {len(parsed)} area values from {selected.name}; skipped {skipped} blank/header/text rows."
            save_calculation(
                operation="Import total",
                input_format="Kanal-Marla-Foot",
                inputs={"file": selected.name, "values": [raw for raw, _ in parsed]},
                settings={"square_feet_per_marla": settings.square_feet_per_marla},
                result=result_text.value,
                details=details_text.value,
            )
            status_text.value = f"Imported and totaled {len(parsed)} values from {selected.name}."
            update_area_previews()
            refresh_history()
        except Exception as exc:
            status_text.value = f"Could not import area values: {exc}"
        page.update()

    import_button = ft.OutlinedButton("Import area list", on_click=import_area_list)

    def clear_history(e=None):
        """Remove saved calculation records after a confirmation prompt."""
        def confirm_clear(_):
            page.pop_dialog()
            try:
                removed = clear_calculations()
                status_text.value = f"Cleared {removed} calculation history records."
                refresh_history()
            except Exception as exc:
                status_text.value = f"Could not clear history: {exc}"
            page.update()

        dialog = ft.AlertDialog(
            modal=True,
            title=ft.Text("Clear calculation history?"),
            content=ft.Text("This removes all saved calculator history."),
            actions=[
                ft.TextButton("Cancel", on_click=lambda _: page.pop_dialog()),
                ft.TextButton("Clear history", on_click=confirm_clear),
            ],
        )
        page.show_dialog(dialog)

    clear_history_button = ft.TextButton("Clear history", on_click=clear_history)

    def update_marla_controls(e=None):
        custom_marla_field.visible = marla_dropdown.value == "custom"
        update_area_previews()
        page.update()

    def update_operation_controls(e=None):
        operation = operation_dropdown.value or "Convert"
        value_b.visible = operation != "Convert"
        preview_b_card.visible = value_b.visible
        if operation in {"Add", "Subtract", "Average"}:
            value_b.label = "B (area: K, K-M, or K-M-F)"
            value_b.hint_text = "Examples: 2, 2-5, or 2-5-100"
        elif operation == "Percentage Share":
            value_b.label = "B (share or percentage)"
            value_b.hint_text = "Enter a fraction such as 1/5 or a percentage"
        elif operation in {"Multiply", "Divide"}:
            value_b.label = "B (number)"
            value_b.hint_text = "Enter a multiplier or divisor"
        else:
            value_b.label = "B"
        update_area_previews()
        if e is not None:
            page.update()

    marla_dropdown.on_select = update_marla_controls
    operation_dropdown.on_select = update_operation_controls
    value_a.on_change = update_area_previews
    value_b.on_change = update_area_previews

    async def calculate(e=None):
        """Calculate using K-M-F areas and preserve area results in field A."""
        try:
            settings = make_settings()
            display_unit = conversion_dropdown.value or "Kanal-Marla-Foot"
            operation = operation_dropdown.value or "Convert"
            first_raw = (value_a.value or "").strip()
            second_raw = (value_b.value or "").strip()
            area_a = parse_area(first_raw, "A", settings)
            # Keep original text in history so the entered K-M-F precision is
            # preserved and the history serializer receives JSON-safe values.
            original_inputs = {
                "A": first_raw,
                "B": second_raw,
            }
            detail_lines = [
                f"Operation: {operation}",
                "Input format: K, K-M, K-M-F, acres, or kilas",
                f"Result display: {display_unit}",
                f"Square feet per marla: {settings.square_feet_per_marla}",
                "Acre: 8 kanals",
                "Kila: 8 kanals + 18 marlas",
            ]

            if operation == "Convert":
                result_area = area_a
                detail_lines.append(f"Converted A = {first_raw} K-M-F.")
            else:
                if operation == "Add":
                    area_b = parse_area(second_raw, "B", settings)
                    result_area = sum_areas([area_a, area_b])
                    detail_lines.append(f"{first_raw} + {second_raw}")
                elif operation == "Subtract":
                    area_b = parse_area(second_raw, "B", settings)
                    result_area = subtract_areas(area_a, area_b)
                    detail_lines.append(f"{first_raw} - {second_raw}")
                elif operation == "Multiply":
                    multiplier = parse_number(second_raw, "B")
                    result_area = LandArea(area_a.square_feet * multiplier)
                    detail_lines.append(f"{first_raw} K-M-F x {multiplier}")
                elif operation == "Divide":
                    divisor = parse_number(second_raw, "B")
                    if divisor == 0:
                        raise AreaError("B cannot be zero when dividing.")
                    result_area = LandArea(area_a.square_feet / divisor)
                    detail_lines.append(f"{first_raw} K-M-F / {divisor}")
                elif operation == "Average":
                    area_b = parse_area(second_raw, "B", settings)
                    result_area = average_area([area_a, area_b])
                    detail_lines.append(f"({first_raw} + {second_raw}) / 2")
                elif operation == "Percentage Share":
                    percentage = parse_share(second_raw)
                    # Preserve fractional square feet for exact acre/kila
                    # shares (for example, 1/5 of 12 kilas is exactly 2.4).
                    result_area = LandArea(
                        area_a.square_feet * percentage / Decimal("100")
                    )
                    detail_lines.append(
                        f"{second_raw} share of {first_raw} = "
                        f"{_format_decimal(result_area.to_kila(settings))} kilas / "
                        f"{_format_decimal(result_area.to_acres(settings))} acres"
                    )
                else:
                    raise AreaError("Choose a supported operation.")

            result = format_result(result_area, settings, display_unit)
            if operation != "Convert":
                # All area-valued arithmetic results return to normalized K-M-F.
                answer = result_area.format(settings)
                value_a.value = answer
                result_text.value = (
                    f"A = {answer} K-M-F\n\n{result}"
                )
            else:
                result_text.value = result

            details_text.value = "\n".join(detail_lines)
            save_calculation(
                operation=operation,
                input_format="Kanal-Marla-Foot",
                inputs=original_inputs,
                settings={
                    "square_feet_per_marla": settings.square_feet_per_marla,
                    "marla_per_kanal": settings.marla_per_kanal,
                    "kanals_per_acre": settings.kanals_per_acre,
                    "kila_kanal": settings.kila_kanal,
                    "kila_marla": settings.kila_marla,
                    "kila_feet": settings.kila_feet,
                    "display_unit": display_unit,
                },
                result=result_text.value,
                details=details_text.value,
            )
            status_text.value = "Calculation completed and saved to history."
            # Keep the calculated area in A for chaining; B is the next operand.
            value_b.value = ""
            update_area_previews()
            refresh_history()
        except (AreaError, ValueError, TypeError) as exc:
            status_text.value = f"Please check your inputs: {exc}"
        except Exception as exc:
            traceback.print_exc()
            status_text.value = f"Unexpected error: {exc}"
        page.update()
        # Flet focus is async; request it after updating the cleared B field.
        if value_b.visible:
            await value_b.focus()

    def clear_calculator(e=None):
        """Clear operands and results without changing regional settings."""
        value_a.value = ""
        value_b.value = ""
        result_text.value = "Enter K, K-M, or K-M-F in A and choose an operation."
        details_text.value = ""
        status_text.value = ""
        update_area_previews()
        page.update()

    calculate_button.on_click = calculate
    clear_button.on_click = clear_calculator
    # Enter from either operand uses the same validated calculation path as the button.
    value_a.on_submit = calculate
    value_b.on_submit = calculate
    custom_marla_field.on_change = update_area_previews
    update_operation_controls()

    def refresh_history_click(e=None):
        refresh_history()
        page.update()

    refresh_button = ft.TextButton(
        "Refresh history",
        on_click=refresh_history_click,
    )
    refresh_history()

    return ft.Column(
        controls=[
            ft.Text("Land Area Calculator", size=26, weight=ft.FontWeight.BOLD),
            ft.Text(
                "Enter K-M-F values or labeled amounts such as 12 kilas or 3 acres. "
                "Convert areas, calculate shares, and continue arithmetic with A.",
                color="#64748B",
            ),
            ft.Container(
                padding=18,
                bgcolor="#FFFFFF",
                border=ft.Border.all(1, "#E2E8F0"),
                border_radius=12,
                content=ft.Column(
                    controls=[
                        ft.Text(
                            "Measurement settings",
                            size=18,
                            weight=ft.FontWeight.BOLD,
                        ),
                        ft.ResponsiveRow(
                            controls=[marla_dropdown, custom_marla_field],
                            run_spacing=8,
                        ),
                        ft.Text(
                            "Default: 272 sq ft per marla; 20 marlas per kanal; "
                            "8 kanals per acre · 8 kanals + 18 marlas per kila.",
                            size=12,
                            color="#64748B",
                        ),
                    ],
                    spacing=10,
                ),
            ),
            ft.Container(
                padding=18,
                bgcolor="#FFFFFF",
                border=ft.Border.all(1, "#E2E8F0"),
                border_radius=12,
                content=ft.Column(
                    controls=[
                        ft.Text(
                            "Calculator",
                            size=18,
                            weight=ft.FontWeight.BOLD,
                        ),
                        ft.ResponsiveRow(
                            controls=[conversion_dropdown, operation_dropdown],
                            run_spacing=8,
                        ),
                        ft.Text(
                            "Enter K-M-F values or labeled amounts such as 12 kilas or 3 acres. "
                            "For multiplication/division B is numeric; shares accept fractions such as 1/5.",
                            size=12,
                            color="#64748B",
                        ),
                        ft.ResponsiveRow(
                            controls=[value_a, value_b],
                            run_spacing=8,
                        ),
                        ft.ResponsiveRow(
                            controls=[preview_a_card, preview_b_card],
                            run_spacing=8,
                        ),
                        ft.Row(
                            controls=[calculate_button, clear_button],
                            spacing=8,
                        ),
                        status_text,
                    ],
                    spacing=12,
                ),
            ),
            ft.Container(
                padding=18,
                bgcolor="#F8FAFC",
                border=ft.Border.all(1, "#E2E8F0"),
                border_radius=12,
                content=ft.Column(
                    controls=[
                        ft.Text("Result", size=18, weight=ft.FontWeight.BOLD),
                        result_text,
                        ft.Divider(),
                        ft.Text("Calculation details", weight=ft.FontWeight.BOLD),
                        details_text,
                    ],
                    spacing=8,
                ),
            ),
            ft.Row(
                controls=[
                    ft.Text(
                        "Calculation history",
                        size=18,
                        weight=ft.FontWeight.BOLD,
                        expand=True,
                    ),
                    refresh_button,
                    clear_history_button,
                ]
            ),
            ft.Row(controls=[import_button], spacing=8),
            export_buttons,
            history_column,
        ],
        spacing=14,
        scroll=ft.ScrollMode.AUTO,
        expand=True,
    )
